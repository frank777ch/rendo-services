"""Genera una copia LIMPIA de las hojas de contabilidad (agua o luz), llenando las
celdas "según recibo" con la API y conservando fórmulas y resultados.

Uso:
    python scripts/limpiar_hojas.py agua  original_agua.xlsx  salida_agua.xlsx  5998198
    python scripts/limpiar_hojas.py luz   original_luz.xlsx   salida_luz.xlsx   1491533

Requisitos: openpyxl, y el .env con credenciales (usa rendo.providers).
Regla: NO se mueven celdas de cálculo ni se cambian fórmulas; solo se colorea/rotula y
se llenan las celdas de entrada del recibo. Ver docs/contabilidad-hojas.md.
"""
from __future__ import annotations

import datetime
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from rendo.config import get_settings
from rendo.providers import build_providers

MESES = {"ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6,
         "JULIO": 7, "AGOSTO": 8, "SETIEMBRE": 9, "SEPTIEMBRE": 9, "OCTUBRE": 10,
         "NOVIEMBRE": 11, "DICIEMBRE": 12}
AMAR = PatternFill("solid", fgColor="FFF2CC")   # lo llena la API
AZUL = PatternFill("solid", fgColor="DDEBF7")   # consumo de inquilinos (manual)
_thin = Side(style="thin", color="BFBFBF")
BORDER = Border(_thin, _thin, _thin, _thin)
MONEY = "#,##0.00"
LEYENDA = ("🟨 Amarillo = lo llena la API (según el recibo)      "
           "🟦 Azul = lo llenas tú (consumo de inquilinos)      ⬜ Blanco = fórmula (no tocar)")


def _dt(s):
    if not s:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def _concepto(r, nombre):
    return next((c["monto"] for c in r.get("conceptos", []) if c["descripcion"] == nombre), None)


def _estilo(ws, titulo, header_color, api_cells, tenant_cells, max_row, legend_row, widths):
    ws["B1"] = titulo
    ws["B1"].font = Font(b=True, size=14, color="FFFFFF")
    ws["B1"].fill = PatternFill("solid", fgColor=header_color)
    ws["B1"].alignment = Alignment(vertical="center")
    try:
        ws.merge_cells("B1:S1")
    except Exception:
        pass
    ws.row_dimensions[1].height = 26
    ws[f"B{legend_row}"] = LEYENDA
    ws[f"B{legend_row}"].font = Font(size=9, italic=True)
    for a in api_cells:
        ws[a].fill = AMAR
    for a in tenant_cells:
        ws[a].fill = AZUL
    for row in ws.iter_rows(min_row=2, max_row=max_row):
        for c in row:
            if c.value is not None:
                c.border = BORDER
                if isinstance(c.value, (int, float)):
                    c.number_format = MONEY
    ws.freeze_panes = "B3"
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def limpiar_agua(src, out, nis):
    P = build_providers(get_settings())["sedapal"]
    wb = openpyxl.load_workbook(src, data_only=False)
    widths = {"B": 24, "C": 12, "D": 11, "E": 13, "F": 11, "G": 13, "H": 12, "I": 12,
              "J": 11, "K": 12, "L": 11, "M": 12, "N": 13, "O": 11, "P": 11, "Q": 11, "R": 13}
    filas = {"0 a 10": 4, "10 a 20": 5, "20 a 50": 6, "50 a más": 7, "50 a mas": 7}
    for name, mn in MESES.items():
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        try:
            r = P.recibo(str(nis), f"2026-{mn:02d}")
        except Exception:
            r = None
        if r:
            if r.get("fecha_tarifa"):
                ws["B2"] = f"Estructura Tarifaria ({r['fecha_tarifa']})"
            for t in (r.get("tarifa") or []):
                row = filas.get(t["rango"])
                if row:
                    ws[f"D{row}"] = t["agua"]
                    ws[f"E{row}"] = t["alcantarillado"]
            if r.get("fecha_emision"):
                ws["G3"] = _dt(r["fecha_emision"])
            if r.get("periodo_inicio"):
                ws["H3"] = _dt(r["periodo_inicio"])
            if r.get("periodo_fin"):
                ws["I3"] = _dt(r["periodo_fin"])
            if r.get("lectura_anterior") is not None:
                ws["H4"] = r["lectura_anterior"]
            if r.get("lectura_actual") is not None:
                ws["I4"] = r["lectura_actual"]
            cf = _concepto(r, "Cargo Fijo")
            if cf is not None:
                ws["H16"] = cf
        api = ["B2", "G3", "H3", "I3", "H4", "I4", "H16"] + [f"{c}{row}" for row in (4, 5, 6, 7) for c in "DE"]
        tenant = [f"M{row}" for row in (10, 11, 12, 13)]
        _estilo(ws, f"AGUA POTABLE  ·  UPIS ROSA  ·  {name} 2026", "1F4E78", api, tenant, 35, 37, widths)
    wb.save(out)


def limpiar_luz(src, out, sum_):
    P = build_providers(get_settings())["luzdelsur"]
    wb = openpyxl.load_workbook(src, data_only=False)
    widths = {"B": 22, "C": 12, "D": 4, "E": 14, "F": 16, "G": 14, "H": 12, "I": 16, "J": 12, "K": 13}
    campos = [("C10", "Cargo Fijo"), ("C11", "Mant. y Reposición de Conexión"),
              ("C12", "Alumbrado Público"), ("C13", "Interés Compensatorio"),
              ("C23", "Electrificación Rural (Ley N° 28749)")]
    for name, mn in MESES.items():
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        try:
            r = P.recibo(str(sum_), f"2026-{mn:02d}")
        except Exception:
            r = None
        if r:
            if r.get("importe_total") is not None:
                ws["J4"] = round(r["importe_total"], 2)
            if r.get("precio_unitario") is not None:
                ws["C5"] = r["precio_unitario"]
            if r.get("lectura_actual") is not None:
                ws["E6"] = r["lectura_actual"]
            if r.get("lectura_anterior") is not None:
                ws["F6"] = r["lectura_anterior"]
            if r.get("fecha_lectura_actual"):
                ws["E4"] = _dt(r["fecha_lectura_actual"])
            if r.get("fecha_lectura_anterior"):
                ws["F4"] = _dt(r["fecha_lectura_anterior"])
            for cell, nom in campos:
                v = _concepto(r, nom)
                if v is not None:
                    ws[cell] = v
        api = ["E4", "F4", "J4", "C5", "E6", "F6", "C10", "C11", "C12", "C13", "C23"]
        tenant = [f"F{row}" for row in (11, 12, 13, 14)]
        _estilo(ws, f"ELECTRICIDAD (LUZ DEL SUR)  ·  UPIS ROSA  ·  {name} 2026", "7F6000", api, tenant, 29, 31, widths)
    wb.save(out)


if __name__ == "__main__":
    if len(sys.argv) != 5:
        print(__doc__)
        sys.exit(1)
    servicio, src, out, sumi = sys.argv[1:5]
    (limpiar_agua if servicio == "agua" else limpiar_luz)(src, out, sumi)
    print(f"generado {out}")
